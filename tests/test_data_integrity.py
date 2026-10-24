"""Fail if a chart value is not in an official source snapshot."""

from __future__ import annotations

import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
SITE = ROOT / "frontend" / "data" / "site.json"

from scripts.fetch.common import number_text, sha256_file  # noqa: E402


def load(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def q(value: Decimal, places: str) -> str:
    return format(value.quantize(Decimal(places), rounding=ROUND_HALF_UP), "f")


def test_manifest_hashes_match_raw_files():
    manifest = load(DATA / "manifest.json")
    assert manifest["sources"], "manifest has no sources"
    for source in manifest["sources"]:
        path = ROOT / source["path"]
        assert path.is_file(), source["path"]
        assert sha256_file(path) == source["sha256"], source["id"]
        assert source["url"].startswith("https://")
        assert source["retrieved_at"]
        assert source["publisher"]


def test_every_chart_series_traces_to_a_source_record():
    site = load(SITE)
    manifest = {source["id"]: source for source in load(DATA / "manifest.json")["sources"]}
    indexes = {
        "bls": _bls_index(manifest),
        "treasury": _treasury_index(manifest),
        "workbooks": {},
        "eia_total": {},
        "eia_intl": {},
        "worldbank": {},
    }

    seen = set()
    for chart in site["charts"]:
        assert chart["series"], chart["id"]
        for series_id in chart["series"]:
            seen.add(series_id)
            _assert_series_matches_source(site["series"][series_id], manifest, indexes)
    for series in site["series"].values():
        if series.get("derived") or series["id"] in seen:
            continue
        _assert_series_matches_source(series, manifest, indexes)


def test_derived_values_recompute_from_source_series():
    site = load(SITE)
    series = site["series"]
    yoy = series["cpi_u_all_items_12_month_change"]
    monthly = {point["date"]: Decimal(point["value"]) for point in series["bls_cuur0000sa0_monthly"]["points"]}
    expected = []
    for date in sorted(monthly):
        year, month = date.split("-")
        prior = f"{int(year) - 1}-{month}"
        if prior not in monthly or monthly[prior] == 0:
            continue
        change = (monthly[date] / monthly[prior] - Decimal(1)) * Decimal(100)
        expected.append((date, q(change, "0.001")))
    assert [(point["date"], point["value"]) for point in yoy["points"]] == expected

    base_year = series["bls_cuur0000saf1_annual_rebased"]["base_year"]
    for source_id, derived_id in (
        ("bls_cuur0000saf1_annual", "bls_cuur0000saf1_annual_rebased"),
        ("bls_cuur0000sah1_annual", "bls_cuur0000sah1_annual_rebased"),
        ("bls_cuur0000sa0e_annual", "bls_cuur0000sa0e_annual_rebased"),
        ("bls_cuur0000sam_annual", "bls_cuur0000sam_annual_rebased"),
        ("bls_cuur0000sae1_annual", "bls_cuur0000sae1_annual_rebased"),
    ):
        raw = {point["date"]: Decimal(point["value"]) for point in series[source_id]["points"]}
        base = raw[base_year]
        got = [(point["date"], point["value"]) for point in series[derived_id]["points"]]
        want = [(year, q(raw[year] / base * Decimal(100), "0.001")) for year in sorted(raw) if year >= base_year]
        assert got == want, derived_id

    prices = {point["date"]: Decimal(point["value"]) for point in series["census_new_home_median"]["points"]}
    cpi = {point["date"]: Decimal(point["value"]) for point in series["bls_cuur0000sa0_annual"]["points"]}
    got = [(point["date"], point["value"]) for point in series["census_new_home_median_cpi_u"]["points"]]
    want = []
    for year in sorted(prices):
        if year in cpi and cpi[year] != 0:
            want.append((year, q(prices[year] * Decimal(100) / cpi[year], "0.01")))
    assert got == want


def test_site_bundle_matches_normalized_files():
    site = load(SITE)
    charts = load(DATA / "normalized" / "charts.json")["charts"]
    assert site["charts"] == charts
    for path in (DATA / "normalized" / "series").glob("*.json"):
        item = load(path)
        assert site["series"][item["id"]]["points"] == item["points"]


def test_energy_derived_values_recompute():
    series = load(SITE)["series"]
    for item in series.values():
        formula = item.get("formula")
        if not formula:
            continue
        if formula == "us_total_with_small_scale_solar":
            _assert_us_total(series, item)
        elif formula == "us_generation_residual":
            _assert_us_residual(series, item)
        elif formula == "share_of_total":
            _assert_share(series, item)
        elif formula == "kwh_per_person":
            _assert_kwh(series, item)
        else:
            raise AssertionError(formula)


def test_energy_charts_cover_required_countries():
    site = load(SITE)
    ids = {chart["id"] for chart in site["charts"]}
    assert "us-electricity-generation-shares" in ids
    for code in ("usa", "gbr", "deu", "fra", "jpn", "chn", "ind", "can"):
        assert f"electricity-generation-{code}" in ids, code
        assert f"electricity-per-person-{code}" in ids, code
    share = next(chart for chart in site["charts"] if chart["id"] == "us-electricity-generation-shares")
    assert share["shade"] == "calendar"
    assert share.get("type") == "stacked"
    usa = next(chart for chart in site["charts"] if chart["id"] == "electricity-per-person-usa")
    assert usa["shade"] == "calendar"
    uk = next(chart for chart in site["charts"] if chart["id"] == "electricity-per-person-gbr")
    assert uk["shade"] == "none"
    assert uk["title"] == "United Kingdom: electricity generation per person, kWh"


def test_no_placeholder_numbers():
    site = load(SITE)
    banned = {"placeholder", "estimated by", "lorem", "todo", "n/a", "null"}
    for series in site["series"].values():
        for point in series["points"]:
            lowered = point["value"].lower()
            assert all(word not in lowered for word in banned)


def _is_number(value: str) -> bool:
    try:
        Decimal(value)
    except Exception:
        return False
    return True


def _assert_series_matches_source(series: dict, manifest: dict, indexes: dict) -> None:
    assert series["points"], series["id"]
    for point in series["points"]:
        assert _is_number(point["value"]), point
        assert point.get("locator"), (series["id"], point)
    if series.get("derived"):
        return
    source = manifest[series["source_id"]]
    fmt = source["format"]
    if fmt == "bls_tsv":
        for point in series["points"]:
            key = (series["bls_series_id"], point["locator"]["year"], point["locator"]["period"])
            assert indexes["bls"][source["id"]][key] == point["value"], (series["id"], point)
    elif fmt == "treasury_json":
        for point in series["points"]:
            assert _treasury_value(indexes["treasury"][source["id"]], point) == point["value"], (series["id"], point)
    elif fmt == "xlsx":
        for point in series["points"]:
            point_source = manifest[point["locator"].get("source_id") or series["source_id"]]
            assert point_source["format"] == "xlsx", (series["id"], point)
            assert _xlsx_value(indexes["workbooks"], point_source, series, point) == point["value"], (series["id"], point)
    elif fmt == "eia_bulk_zip":
        table = _eia_total_table(indexes, source)
        for point in series["points"]:
            got = table[point["locator"]["series_id"]][point["locator"]["year"]]
            assert Decimal(got) == Decimal(point["value"]), (series["id"], point)
    elif fmt == "eia_intl_jsonl":
        table = _eia_intl_table(indexes, source)
        for point in series["points"]:
            got = table[point["locator"]["series_id"]][point["locator"]["year"]]
            assert Decimal(got) == Decimal(point["value"]), (series["id"], point)
    elif fmt == "worldbank_zip":
        table = _worldbank_table(indexes, source)
        for point in series["points"]:
            got = table[point["locator"]["country_code"]][point["locator"]["year"]]
            assert Decimal(got) == Decimal(point["value"]), (series["id"], point)
    else:
        raise AssertionError(f"unknown source format {fmt}")


def _year_values(series: dict) -> dict[str, Decimal]:
    return {point["date"]: Decimal(point["value"]) for point in series["points"]}


def _assert_us_total(series: dict, item: dict) -> None:
    base_id, solar_id, solar_all_id = item["inputs"]
    base = _year_values(series[base_id])
    solar = _year_values(series[solar_id])
    solar_all = _year_values(series[solar_all_id])
    expected = []
    for year in sorted(base):
        if year in solar and year in solar_all:
            value = base[year] + solar[year] - solar_all[year]
        else:
            value = base[year]
        expected.append((year, q(value, "0.001")))
    assert [(point["date"], point["value"]) for point in item["points"]] == expected


def _assert_us_residual(series: dict, item: dict) -> None:
    total_id, *part_ids = item["inputs"]
    total = _year_values(series[total_id])
    parts = [_year_values(series[part_id]) for part_id in part_ids]
    expected = []
    for year in sorted(total):
        named = sum((part[year] for part in parts if year in part), Decimal(0))
        expected.append((year, q(total[year] - named, "0.001")))
    assert [(point["date"], point["value"]) for point in item["points"]] == expected


def _assert_share(series: dict, item: dict) -> None:
    level_id, total_id = item["inputs"]
    level = _year_values(series[level_id])
    total = _year_values(series[total_id])
    expected = []
    for year in sorted(level):
        if year not in total or total[year] == 0:
            continue
        expected.append((year, q(level[year] / total[year] * Decimal(100), "0.001")))
    assert [(point["date"], point["value"]) for point in item["points"]] == expected


def _assert_kwh(series: dict, item: dict) -> None:
    gen_id, pop_id = item["inputs"]
    generation = _year_values(series[gen_id])
    population = _year_values(series[pop_id])
    expected = []
    for year in sorted(set(generation) & set(population)):
        value = generation[year] * Decimal(1000000000) / population[year]
        expected.append((year, q(value, "1")))
    assert [(point["date"], point["value"]) for point in item["points"]] == expected


def _decimal_token(value) -> str | None:
    if isinstance(value, bool) or value is None or isinstance(value, str):
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    return None


def _eia_total_table(indexes: dict, source: dict) -> dict:
    if source["id"] not in indexes["eia_total"]:
        import zipfile

        table = {}
        with zipfile.ZipFile(ROOT / source["path"]) as archive:
            name = next(item for item in archive.namelist() if item.endswith("TOTAL.txt"))
            with archive.open(name) as handle:
                for raw in handle:
                    line = raw.decode("utf-8").strip()
                    if not line.startswith("{"):
                        continue
                    obj = json.loads(line, parse_float=Decimal)
                    series_id = obj.get("series_id")
                    if not series_id or not isinstance(obj.get("data"), list):
                        continue
                    years = {}
                    for pair in obj["data"]:
                        if not isinstance(pair, list) or len(pair) < 2:
                            continue
                        year, value = pair[0], pair[1]
                        token = _decimal_token(value)
                        if token is None:
                            continue
                        years[str(year)] = token
                    table[series_id] = years
        indexes["eia_total"][source["id"]] = table
    return indexes["eia_total"][source["id"]]


def _eia_intl_table(indexes: dict, source: dict) -> dict:
    if source["id"] not in indexes["eia_intl"]:
        table = {}
        for line in (ROOT / source["path"]).read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            obj = json.loads(line, parse_float=Decimal)
            years = {}
            for year, value in obj.get("data") or []:
                token = _decimal_token(value)
                if token is None:
                    continue
                years[str(year)] = token
            table[obj["series_id"]] = years
        indexes["eia_intl"][source["id"]] = table
    return indexes["eia_intl"][source["id"]]


def _worldbank_table(indexes: dict, source: dict) -> dict:
    if source["id"] not in indexes["worldbank"]:
        import csv
        import zipfile

        with zipfile.ZipFile(ROOT / source["path"]) as archive:
            name = next(item for item in archive.namelist() if item.startswith("API_SP.POP.TOTL") and "Metadata" not in item)
            text = archive.read(name).decode("utf-8-sig")
        lines = text.splitlines()
        start = next(index for index, line in enumerate(lines) if line.startswith('"Country Name"') or line.startswith("Country Name"))
        table = {}
        for row in csv.DictReader(lines[start:]):
            code = (row.get("Country Code") or "").strip()
            years = {}
            for key, value in row.items():
                if key and key.isdigit() and value and number_text(value):
                    years[key] = number_text(value)
            if code:
                table[code] = years
        indexes["worldbank"][source["id"]] = table
    return indexes["worldbank"][source["id"]]


def _bls_index(manifest: dict) -> dict:
    indexed = {}
    for source in manifest.values():
        if source["format"] != "bls_tsv":
            continue
        table = {}
        lines = (ROOT / source["path"]).read_text(encoding="utf-8").splitlines()
        for line in lines[1:]:
            if not line.strip():
                continue
            parts = line.split("\t")
            table[(parts[0].strip(), parts[1].strip(), parts[2].strip())] = parts[3].strip()
        indexed[source["id"]] = table
    return indexed


def _treasury_index(manifest: dict) -> dict:
    indexed = {}
    for source in manifest.values():
        if source["format"] != "treasury_json":
            continue
        indexed[source["id"]] = load(ROOT / source["path"])["data"]
    return indexed


def _treasury_value(rows: list[dict], point: dict) -> str:
    locator = point["locator"]
    if locator.get("aggregate") == "sum":
        total = Decimal("0.00")
        found = False
        for row in rows:
            if row.get("record_date") == locator["record_date"]:
                found = True
                total += Decimal(row[locator["field"]])
        assert found
        return format(total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP), "f")
    matches = [
        row
        for row in rows
        if row.get("record_date") == locator["record_date"]
        and (locator.get("sequence_number_cd") is None or str(row.get("sequence_number_cd")) == str(locator["sequence_number_cd"]))
    ]
    assert len(matches) == 1, locator
    return str(matches[0][locator["field"]])


XLSX_COLUMNS = {
    "omb_receipts": 1,
    "omb_outlays": 2,
    "omb_deficit": 3,
    "omb_gdp": 1,
    "omb_receipts_pct_gdp": 2,
    "omb_outlays_pct_gdp": 3,
    "omb_deficit_pct_gdp": 4,
    "omb_receipts_fy2017": 4,
    "omb_outlays_fy2017": 5,
    "omb_deficit_fy2017": 6,
    "omb_debt_gross": 1,
    "omb_debt_public": 3,
    "omb_debt_gross_pct_gdp": 6,
    "omb_debt_public_pct_gdp": 8,
    "census_median_income_current": 2,
    "census_median_income_adjusted": 3,
    "census_new_home_median": 1,
}


def _xlsx_value(cache: dict, source: dict, series: dict, point: dict) -> str:
    path = str(ROOT / source["path"])
    locator = point["locator"]
    if series["id"] == "dhs_lpr_table1":
        return _lpr_table_value(cache, path, locator["row_label"])
    if locator.get("aggregate") == "sum_column":
        return _sum_column_value(cache, path, locator)
    if path not in cache:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        cache[path] = workbook
    workbook = cache[path]
    if "line_code" in locator:
        sheet = workbook[locator["sheet"]]
        rows = list(sheet.iter_rows(values_only=True))
        header = next(row for row in rows if row and row[0] == "Line")
        data = next(row for row in rows if row and str(row[2]).strip() == locator["line_code"])
        column = next(index for index, cell in enumerate(header) if str(cell).strip() == locator["column"])
        return number_text(data[column])
    if series["id"] == "census_new_home_median":
        sheet = workbook["Annual"]
        for row in sheet.iter_rows(values_only=True):
            if row and str(row[0]).strip() == point["date"]:
                return number_text(row[1])
        raise AssertionError(point)
    if series["source_id"] == "census_h05":
        column = XLSX_COLUMNS[series["id"]]
        for row in workbook[workbook.sheetnames[0]].iter_rows(values_only=True):
            if row and str(row[0]).strip() == locator["row_label"]:
                return number_text(row[column])
        raise AssertionError(locator)
    if series["source_id"] == "hist03z1":
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[1]
        column = next(index for index, cell in enumerate(header) if str(cell).strip() == locator["column"])
        for row in rows:
            label = " ".join(str(row[0]).split()) if row and row[0] is not None else ""
            if label == locator["row_label"] and number_text(row[column]) == point["value"]:
                return point["value"]
        raise AssertionError(locator)
    column = XLSX_COLUMNS[series["id"]]
    for row in workbook[workbook.sheetnames[0]].iter_rows(values_only=True):
        if row and str(row[0]).strip() == point["date"]:
            return number_text(row[column])
    raise AssertionError(point)


def _lpr_table_value(cache: dict, path: str, row_label: str) -> str:
    key = (path, "lpr_table1")
    if key not in cache:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        index: dict[str, str] = {}
        for row in workbook["Table 1"].iter_rows(values_only=True):
            for offset in range(0, len(row) - 1, 2):
                year, number = row[offset], row[offset + 1]
                if isinstance(year, int) and not isinstance(year, bool):
                    label = str(year)
                elif isinstance(year, str):
                    label = " ".join(year.split())
                else:
                    continue
                if isinstance(number, bool) or not isinstance(number, int):
                    continue
                assert label not in index, label
                index[label] = str(number)
        workbook.close()
        cache[key] = index
    assert row_label in cache[key], row_label
    return cache[key][row_label]


def _sum_column_value(cache: dict, path: str, locator: dict) -> str:
    key = (path, "sum_column", locator["sheet"], locator["field"])
    if key not in cache:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        iterator = workbook[locator["sheet"]].iter_rows(values_only=True)
        header = [" ".join(str(cell).split()) if cell is not None else "" for cell in next(iterator)]
        year_index = header.index("Fiscal Year")
        month_index = header.index("Month (abbv)")
        count_index = header.index(locator["field"])
        totals: dict[tuple[str, str], int] = {}
        for row in iterator:
            month = row[month_index].strip().upper() if isinstance(row[month_index], str) else str(row[month_index])
            group = (str(row[year_index]).strip(), month)
            totals[group] = totals.get(group, 0) + int(row[count_index])
        workbook.close()
        cache[key] = totals
    group = (locator["fiscal_year"], locator["month"])
    assert group in cache[key], group
    return str(cache[key][group])
