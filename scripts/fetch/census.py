"""Census income table H-5 and new-house sales prices."""

from __future__ import annotations

import openpyxl

from .common import RAW, download, is_year, number_text, source_record, utc_now, write_bytes

PUBLISHER = "U.S. Census Bureau"

INCOME_URL = "https://www2.census.gov/programs-surveys/cps/tables/time-series/historical-income-households/h05.xlsx"
INCOME_PAGE = "https://www.census.gov/data/tables/time-series/demo/income-poverty/historical-income-households.html"
PRICE_URL = "https://www.census.gov/construction/nrs/xls/price_cust.xlsx"
PRICE_PAGE = "https://www.census.gov/construction/nrs/data/series.html"


def _year_token(value: object) -> tuple[str, bool] | None:
    text = str(value).strip() if value is not None else ""
    if len(text) >= 4 and text[:4].isdigit() and (len(text) == 4 or text[4] == " "):
        return text[:4], "(" in text
    return None


def fetch() -> tuple[list[dict], list[dict]]:
    retrieved = utc_now()
    sources: list[dict] = []
    series: list[dict] = []

    income_payload = download(INCOME_URL, timeout=180)
    income_path = RAW / "census" / "h05.xlsx"
    write_bytes(income_path, income_payload)
    workbook = openpyxl.load_workbook(income_path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    title = " ".join(str(rows[1][0]).split())
    adjustment_note = " ".join(str(rows[5][0]).split())
    sources.append(
        source_record(
            "census_h05",
            PUBLISHER,
            title,
            INCOME_URL,
            income_path,
            retrieved,
            "xlsx",
            {"page_url": INCOME_PAGE, "adjustment_note": adjustment_note},
        )
    )
    # All Races block runs until the next non-year section label.
    chosen: dict[str, tuple] = {}
    in_block = False
    for row in rows:
        label = str(row[0]).strip() if row and row[0] is not None else ""
        if label == "All Races":
            in_block = True
            continue
        if not in_block:
            continue
        if label and not label[0].isdigit() and label not in {"Race, Hispanic origin, and year"}:
            break
        parsed = _year_token(row[0])
        if parsed is None:
            continue
        year, footnoted = parsed
        # Table H-5 lists the row used for comparison with later years first when a year is repeated.
        if year not in chosen:
            chosen[year] = (row, footnoted, label)
    current_points = []
    adjusted_points = []
    for year in sorted(chosen):
        row, _footnoted, label = chosen[year]
        current = number_text(row[2] if len(row) > 2 else None)
        adjusted = number_text(row[3] if len(row) > 3 else None)
        if current is not None:
            current_points.append(
                {"date": year, "value": current, "locator": {"row_label": label, "column": "Current dollars"}}
            )
        if adjusted is not None:
            adjusted_points.append(
                {"date": year, "value": adjusted, "locator": {"row_label": label, "column": "2025 dollars"}}
            )
    series.append(
        {
            "id": "census_median_income_current",
            "name": "Median income, Current dollars",
            "units": "Current dollars",
            "adjustment": "not inflation-adjusted",
            "nominal_or_real": "nominal",
            "frequency": "annual",
            "source_id": "census_h05",
            "series_url": INCOME_PAGE,
            "notes": adjustment_note + " All Races. When a year is listed twice, the first row is used.",
            "points": current_points,
        }
    )
    series.append(
        {
            "id": "census_median_income_adjusted",
            "name": "Median income, 2025 dollars",
            "units": "2025 dollars",
            "adjustment": adjustment_note,
            "nominal_or_real": "inflation-adjusted",
            "frequency": "annual",
            "source_id": "census_h05",
            "series_url": INCOME_PAGE,
            "notes": "Inflation adjustment is the one Census publishes in this column, not a CPI-U adjustment computed by this site. " + adjustment_note,
            "points": adjusted_points,
        }
    )

    price_payload = download(PRICE_URL, timeout=180)
    price_path = RAW / "census" / "price_cust.xlsx"
    write_bytes(price_path, price_payload)
    workbook = openpyxl.load_workbook(price_path, data_only=True, read_only=True)
    sheet = workbook["Annual"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    title = " ".join(str(rows[0][0]).split())
    land_note = ""
    for row in rows:
        text = " ".join(str(row[0]).split()) if row and row[0] else ""
        if text.lower().startswith("note:"):
            land_note = text
    sources.append(
        source_record(
            "census_new_home_prices",
            PUBLISHER,
            title,
            PRICE_URL,
            price_path,
            retrieved,
            "xlsx",
            {"page_url": PRICE_PAGE, "note": land_note},
        )
    )
    points = []
    for row in rows:
        if not row or not is_year(row[0]):
            continue
        value = number_text(row[1] if len(row) > 1 else None)
        if value is None:
            continue
        points.append(
            {
                "date": str(row[0]).strip(),
                "value": value,
                "locator": {"sheet": "Annual", "column": "Median sales price", "year": str(row[0]).strip()},
            }
        )
    series.append(
        {
            "id": "census_new_home_median",
            "name": "Median sales price",
            "units": "Dollars, nominal",
            "adjustment": "not inflation-adjusted",
            "nominal_or_real": "nominal",
            "frequency": "annual",
            "source_id": "census_new_home_prices",
            "series_url": PRICE_PAGE,
            "notes": f"{title}. United States. {land_note} New privately owned houses sold. This is not the price of existing homes.",
            "points": points,
        }
    )
    return sources, series
