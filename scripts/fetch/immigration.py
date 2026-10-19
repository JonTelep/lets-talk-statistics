"""DHS lawful permanent residents and CBP nationwide encounters.

A chart is written only when the workbook layout matches the parser below.
Anything else stays a TODO. No counts are filled in.
"""

from __future__ import annotations

import re
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

import openpyxl

from .common import RAW, download, source_record, utc_now, write_bytes

YEARBOOK_INDEX = "https://ohss.dhs.gov/topics/immigration/yearbook"
CBP_PAGE = "https://www.cbp.gov/document/stats/cbp-data-download-portal-library"

DHS_DATASET = "DHS Yearbook of Immigration Statistics, Table 1, Persons Obtaining Lawful Permanent Resident Status"
CBP_DATASET = "CBP nationwide encounters"

OMITTED = "The page downloaded, but this run did not find a workbook link it could parse without guessing. The chart is omitted."

MONTH_TO_CALENDAR = {
    "OCT": (-1, 10),
    "NOV": (-1, 11),
    "DEC": (-1, 12),
    "JAN": (0, 1),
    "FEB": (0, 2),
    "MAR": (0, 3),
    "APR": (0, 4),
    "MAY": (0, 5),
    "JUN": (0, 6),
    "JUL": (0, 7),
    "AUG": (0, 8),
    "SEP": (0, 9),
}


class _LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._href: str | None = None
        self._chunks: list[str] = []
        self.links: list[tuple[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self._href = dict(attrs).get("href")
            self._chunks = []

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self._href is not None:
            self.links.append((self._href, " ".join("".join(self._chunks).split())))
            self._href = None

    def handle_data(self, data: str) -> None:
        if self._href is not None:
            self._chunks.append(data)


def _links(html: str) -> list[tuple[str, str]]:
    parser = _LinkParser()
    parser.feed(html)
    return parser.links


def _todo(item_id: str, dataset: str, url: str, reason: str) -> dict:
    return {"id": item_id, "dataset": dataset, "url": url, "reason": reason}


def _plain(text: str) -> str:
    return (
        text.replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2015", "-")
        .replace("\u00a0", " ")
    )


def _slug(filename: str) -> str:
    stem = Path(filename).stem.lower()
    return re.sub(r"[^a-z0-9]+", "_", stem).strip("_")


def _filename(url: str) -> str:
    return unquote(Path(urlparse(url).path).name)


def _year_label(value: object) -> str | None:
    if isinstance(value, int) and not isinstance(value, bool) and 1000 <= value <= 9999:
        return str(value)
    if isinstance(value, str):
        text = " ".join(value.split())
        match = re.fullmatch(r"(\d{4})(?: (\d+))?", text)
        if match:
            return text
    return None


def _table_pairs(row: tuple) -> list[tuple[str, int]] | None:
    found: list[tuple[str, int]] = []
    saw_value = False
    for index in range(0, len(row) - 1, 2):
        year, number = row[index], row[index + 1]
        if year is None and number is None:
            continue
        saw_value = True
        label = _year_label(year)
        if label is None or isinstance(number, bool) or not isinstance(number, int):
            return None
        found.append((label, number))
    if not saw_value:
        return []
    return found


def _parse_lpr_table(payload: bytes) -> tuple[list[dict], str]:
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=True, read_only=True)
    try:
        if "Table 1" not in workbook.sheetnames:
            raise ValueError("Table 1 sheet is missing")
        rows = [tuple(row) for row in workbook["Table 1"].iter_rows(values_only=True)]
    finally:
        workbook.close()

    title = " ".join(str(cell) for row in rows[:6] for cell in row if cell is not None)
    if "LAWFUL PERMANENT RESIDENT" not in title.upper():
        raise ValueError("Table 1 title is not lawful permanent resident status")

    header_at = None
    for index, row in enumerate(rows):
        cells = [" ".join(str(cell).split()) if cell is not None else "" for cell in row]
        if cells and cells[0] == "Year" and len(cells) > 1 and cells[1] == "Number":
            header_at = index
            break
    if header_at is None:
        raise ValueError("Table 1 year header is missing")

    points: list[dict] = []
    seen: set[str] = set()
    footnotes: list[str] = []
    started = False
    for row in rows[header_at + 1 :]:
        parsed = _table_pairs(row)
        if parsed is None:
            text = _plain(" ".join(" ".join(str(cell).split()) for cell in row if cell is not None))
            if not started:
                raise ValueError("Table 1 has a row that is not a year and a count")
            if text:
                footnotes.append(text)
            continue
        if not parsed:
            continue
        started = True
        for label, number in parsed:
            if label in seen:
                raise ValueError(f"Table 1 repeats {label}")
            seen.add(label)
            points.append(
                {
                    "date": label[:4],
                    "value": str(number),
                    "locator": {
                        "source_id": "dhs_lpr_yearbook",
                        "sheet": "Table 1",
                        "row_label": label,
                        "column": "Number",
                    },
                }
            )
    if len(points) < 2:
        raise ValueError("Table 1 did not yield a series")
    points.sort(key=lambda point: point["date"])
    return points, " ".join(footnotes)


def _parse_encounter_counts(payload: bytes) -> dict[tuple[int, str], int]:
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=True, read_only=True)
    try:
        if "Data" not in workbook.sheetnames:
            raise ValueError("Data sheet is missing")
        iterator = workbook["Data"].iter_rows(values_only=True)
        header = [" ".join(str(cell).split()) if cell is not None else "" for cell in next(iterator)]
        needed = ("Fiscal Year", "Month (abbv)", "Encounter Count")
        if any(name not in header for name in needed):
            raise ValueError("Encounter Count columns are missing")
        year_index = header.index("Fiscal Year")
        month_index = header.index("Month (abbv)")
        count_index = header.index("Encounter Count")
        totals: dict[tuple[int, str], int] = {}
        for row in iterator:
            year = row[year_index]
            month = row[month_index].strip().upper() if isinstance(row[month_index], str) else row[month_index]
            count = row[count_index]
            if (
                isinstance(year, bool)
                or not isinstance(year, int)
                or month not in MONTH_TO_CALENDAR
                or isinstance(count, bool)
                or not isinstance(count, int)
                or count < 0
            ):
                raise ValueError("A row is not a fiscal year, month, and count")
            key = (year, month)
            totals[key] = totals.get(key, 0) + count
    finally:
        workbook.close()
    if not totals:
        raise ValueError("Workbook has no encounter rows")
    return totals


def _fetch_dhs(retrieved: str) -> tuple[list[dict], list[dict], list[dict]]:
    try:
        index_html = download(YEARBOOK_INDEX, timeout=60).decode("utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        return [], [], [_todo("dhs", DHS_DATASET, YEARBOOK_INDEX, f"Download failed ({exc.__class__.__name__}). The chart is omitted.")]

    years = []
    for href, _text in _links(index_html):
        match = re.search(r"/topics/immigration/yearbook/(\d{4})$", urljoin(YEARBOOK_INDEX, href))
        if match:
            years.append(int(match.group(1)))
    if not years:
        return [], [], [_todo("dhs", DHS_DATASET, YEARBOOK_INDEX, OMITTED)]
    page_url = f"https://ohss.dhs.gov/topics/immigration/yearbook/{max(years)}"

    try:
        page_html = download(page_url, timeout=60).decode("utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        return [], [], [_todo("dhs", DHS_DATASET, page_url, f"Download failed ({exc.__class__.__name__}). The chart is omitted.")]

    workbooks: dict[str, str] = {}
    for href, text in _links(page_html):
        label = _plain(text)
        lowered = label.lower()
        if "lawful permanent resident" not in lowered or "data table" not in lowered:
            continue
        if not href.lower().split("?", 1)[0].endswith(".xlsx"):
            continue
        workbooks.setdefault(urljoin(page_url, href), label)
    if len(workbooks) != 1:
        return [], [], [_todo("dhs", DHS_DATASET, page_url, OMITTED)]

    file_url, label = next(iter(workbooks.items()))
    try:
        payload = download(file_url, timeout=180)
        points, footnotes = _parse_lpr_table(payload)
    except Exception as exc:  # noqa: BLE001
        detail = _plain(str(exc)).strip() or exc.__class__.__name__
        return [], [], [_todo("dhs", DHS_DATASET, page_url, f"The workbook downloaded, but this run could not map Table 1 without guessing ({detail}). The chart is omitted.")]

    path = RAW / "dhs" / _filename(file_url)
    write_bytes(path, payload)
    source = source_record(
        "dhs_lpr_yearbook",
        "Office of Homeland Security Statistics",
        label,
        file_url,
        path,
        retrieved,
        "xlsx",
        {"page_url": page_url, "series_url": f"{page_url}/table1"},
    )
    note = "Table 1, Persons Obtaining Lawful Permanent Resident Status."
    if footnotes:
        note = f"{note} {footnotes}"
    series = {
        "id": "dhs_lpr_table1",
        "name": "Persons Obtaining Lawful Permanent Resident Status",
        "units": "Persons",
        "adjustment": "as published",
        "nominal_or_real": "count",
        "frequency": "fiscal_year",
        "source_id": "dhs_lpr_yearbook",
        "series_url": f"{page_url}/table1",
        "notes": note,
        "points": points,
    }
    return [source], [series], []


def _fetch_cbp(retrieved: str) -> tuple[list[dict], list[dict], list[dict]]:
    try:
        page_html = download(CBP_PAGE, timeout=60).decode("utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        return [], [], [_todo("cbp", CBP_DATASET, CBP_PAGE, f"Download failed ({exc.__class__.__name__}). The chart is omitted.")]

    books: list[tuple[str, str]] = []
    seen: set[str] = set()
    for href, text in _links(page_html):
        label = _plain(text)
        if not label.lower().startswith("nationwide encounters"):
            continue
        if not href.lower().split("?", 1)[0].endswith(".xlsx"):
            continue
        file_url = urljoin(CBP_PAGE, href)
        if file_url in seen:
            continue
        seen.add(file_url)
        books.append((file_url, label))
    if not books:
        return [], [], [_todo("cbp", CBP_DATASET, CBP_PAGE, OMITTED)]

    parsed: list[tuple[str, str, str, bytes, dict[tuple[int, str], int]]] = []
    combined: dict[tuple[int, str], str] = {}
    for file_url, label in books:
        try:
            payload = download(file_url, timeout=180)
            totals = _parse_encounter_counts(payload)
        except Exception as exc:  # noqa: BLE001
            detail = _plain(str(exc)).strip() or exc.__class__.__name__
            reason = (
                f"The workbook downloaded ({label}), but this run could not sum Encounter Count "
                f"without guessing ({detail}). The chart is omitted."
            )
            return [], [], [_todo("cbp", CBP_DATASET, CBP_PAGE, reason)]
        source_id = "cbp_" + _slug(_filename(file_url))
        for key in totals:
            if key in combined:
                reason = (
                    "Two nationwide encounters workbooks list the same fiscal year and month. "
                    "The chart is omitted rather than choosing a file."
                )
                return [], [], [_todo("cbp", CBP_DATASET, CBP_PAGE, reason)]
            combined[key] = source_id
        parsed.append((source_id, file_url, label, payload, totals))

    sources = []
    for source_id, file_url, label, payload, _totals in parsed:
        path = RAW / "cbp" / _filename(file_url)
        write_bytes(path, payload)
        sources.append(
            source_record(
                source_id,
                "U.S. Customs and Border Protection",
                label,
                file_url,
                path,
                retrieved,
                "xlsx",
                {"page_url": CBP_PAGE, "series_url": CBP_PAGE},
            )
        )

    points = []
    for year, month in sorted(combined, key=lambda item: _calendar_date(item[0], item[1])):
        offset, month_number = MONTH_TO_CALENDAR[month]
        source_id = combined[(year, month)]
        total = next(item[4][(year, month)] for item in parsed if item[0] == source_id)
        points.append(
            {
                "date": f"{year + offset:04d}-{month_number:02d}",
                "value": str(total),
                "locator": {
                    "source_id": source_id,
                    "sheet": "Data",
                    "aggregate": "sum_column",
                    "field": "Encounter Count",
                    "fiscal_year": str(year),
                    "month": month,
                },
            }
        )
    labels = "; ".join(label for _url, label in books)
    series = {
        "id": "cbp_nationwide_encounters_monthly",
        "name": "Nationwide encounters",
        "units": "Encounters",
        "adjustment": "as published",
        "nominal_or_real": "count",
        "frequency": "monthly",
        "source_id": points[-1]["locator"]["source_id"],
        "series_url": CBP_PAGE,
        "notes": (
            "Sum of the Encounter Count column for each fiscal year and month. "
            f"Workbooks: {labels}. "
            "October of a fiscal year is October of the previous calendar year. "
            "Repeated rows stay in the sum."
        ),
        "points": points,
    }
    return sources, [series], []


def _calendar_date(year: int, month: str) -> str:
    offset, month_number = MONTH_TO_CALENDAR[month]
    return f"{year + offset:04d}-{month_number:02d}"


def fetch() -> tuple[list[dict], list[dict], list[dict]]:
    retrieved = utc_now()
    sources: list[dict] = []
    series: list[dict] = []
    todos: list[dict] = []
    for loader in (_fetch_dhs, _fetch_cbp):
        got_sources, got_series, got_todos = loader(retrieved)
        sources.extend(got_sources)
        series.extend(got_series)
        todos.extend(got_todos)
    return sources, series, todos
