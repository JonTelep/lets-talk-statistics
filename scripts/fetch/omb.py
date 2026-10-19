"""OMB Historical Tables, Budget of the United States Government, FY 2027."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from .common import RAW, download, is_year, number_text, source_record, utc_now, write_bytes

PUBLISHER = "Office of Management and Budget"
PAGE = "https://www.whitehouse.gov/omb/information-resources/budget/historical-tables/"
BASE = "https://www.whitehouse.gov/wp-content/uploads/2026/04/"

TABLES = {
    "hist01z1_fy2027.xlsx": "hist01z1",
    "hist01z2_fy2027.xlsx": "hist01z2",
    "hist01z3_fy2027.xlsx": "hist01z3",
    "hist03z1_fy2027.xlsx": "hist03z1",
    "hist07z1_fy2027.xlsx": "hist07z1",
}


def _rows(path: Path) -> list[tuple]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def _cell(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split())


def _year_points(rows: list[tuple], value_index: int, label: str) -> list[dict]:
    points = []
    for row in rows:
        if not row or not is_year(row[0]):
            continue
        value = number_text(row[value_index] if value_index < len(row) else None)
        if value is None:
            continue
        points.append(
            {
                "date": str(row[0]).strip(),
                "value": value,
                "locator": {"row_label": str(row[0]).strip(), "column": label},
            }
        )
    return points


def _series(series_id: str, name: str, units: str, nominal_or_real: str, source_id: str, url: str, points: list[dict], notes: str) -> dict:
    return {
        "id": series_id,
        "name": name,
        "units": units,
        "adjustment": "fiscal year",
        "nominal_or_real": nominal_or_real,
        "frequency": "fiscal_year",
        "source_id": source_id,
        "series_url": url,
        "notes": notes,
        "points": points,
    }


def fetch() -> tuple[list[dict], list[dict]]:
    retrieved = utc_now()
    sources: list[dict] = []
    series: list[dict] = []
    paths: dict[str, Path] = {}
    for filename, source_id in TABLES.items():
        url = BASE + filename
        payload = download(url, timeout=180)
        path = RAW / "omb" / filename
        write_bytes(path, payload)
        paths[source_id] = path
        sources.append(
            source_record(
                source_id,
                PUBLISHER,
                filename,
                url,
                path,
                retrieved,
                "xlsx",
                {"page_url": PAGE},
            )
        )

    rows = _rows(paths["hist01z1"])
    title = _cell(rows[0][0])
    units = _cell(rows[1][0]) or "millions of dollars"
    note = ""
    for row in rows:
        text = _cell(row[0]) if row else ""
        if text.lower().startswith("note:"):
            note = text
    sources[0]["title"] = title
    url = BASE + "hist01z1_fy2027.xlsx"
    for index, (series_id, name) in enumerate(
        [
            ("omb_receipts", "Receipts"),
            ("omb_outlays", "Outlays"),
            ("omb_deficit", "Surplus or Deficit (-)"),
        ],
        start=1,
    ):
        series.append(
            _series(
                series_id,
                name,
                units,
                "nominal",
                "hist01z1",
                url,
                _year_points(rows, index, name),
                note,
            )
        )

    rows = _rows(paths["hist01z2"])
    title = _cell(rows[0][0])
    sources[1]["title"] = title
    url = BASE + "hist01z2_fy2027.xlsx"
    # Row 2 labels: GDP, Receipts, Outlays, Surplus or Deficit (-)
    for index, (series_id, name, kind) in enumerate(
        [
            ("omb_gdp", "GDP (in billions of dollars)", "nominal"),
            ("omb_receipts_pct_gdp", "Receipts", "share_of_gdp"),
            ("omb_outlays_pct_gdp", "Outlays", "share_of_gdp"),
            ("omb_deficit_pct_gdp", "Surplus or Deficit (-)", "share_of_gdp"),
        ],
        start=1,
    ):
        units = "billions of dollars" if series_id == "omb_gdp" else "percent of GDP"
        series.append(
            _series(series_id, name, units, kind, "hist01z2", url, _year_points(rows, index, name), title)
        )

    rows = _rows(paths["hist01z3"])
    title = _cell(rows[0][0])
    sources[2]["title"] = title
    url = BASE + "hist01z3_fy2027.xlsx"
    # Constant FY 2017 dollars are columns 4, 5, 6 under the header in row 2.
    header = " ".join(_cell(cell) for cell in rows[2])
    if "FY 2017" not in header and "FY 2017" not in title:
        raise RuntimeError("Table 1.3 constant-dollar base was not FY 2017")
    for index, (series_id, name) in enumerate(
        [
            ("omb_receipts_fy2017", "Receipts"),
            ("omb_outlays_fy2017", "Outlays"),
            ("omb_deficit_fy2017", "Surplus or Deficit (-)"),
        ],
        start=4,
    ):
        series.append(
            _series(
                series_id,
                name,
                "billions of constant FY 2017 dollars",
                "inflation-adjusted",
                "hist01z3",
                url,
                _year_points(rows, index, name),
                "Constant dollars use the composite deflator published in this table. Base: FY 2017 dollars.",
            )
        )

    rows = _rows(paths["hist03z1"])
    title = _cell(rows[0][0])
    sources[3]["title"] = title
    url = BASE + "hist03z1_fy2027.xlsx"
    header = rows[1]
    year_columns = []
    for index, cell in enumerate(header):
        text = _cell(cell)
        if is_year(text):
            year_columns.append((index, text))
        elif "estimate" in text.lower():
            continue
    dollar_labels = {
        "National Defense": "omb_fn_defense",
        "Human resources": "omb_fn_human",
        "Physical resources": "omb_fn_physical",
        "Net interest": "omb_fn_interest",
        "Other functions": "omb_fn_other",
        "Total, Federal outlays": "omb_fn_total",
    }
    by_label = {}
    section = "dollars"
    for row in rows[2:]:
        label = _cell(row[0]) if row else ""
        if label.lower().startswith("as percentages of outlays"):
            section = "outlays_pct"
            continue
        if label.lower().startswith("as percentages of gdp"):
            section = "gdp_pct"
            continue
        if section == "dollars" and label in dollar_labels:
            by_label[label] = row
        if section == "gdp_pct" and label.lower() in {item.lower() for item in dollar_labels}:
            by_label[label.lower() + "|gdp"] = row
    for label, series_id in dollar_labels.items():
        row = by_label.get(label)
        if row is None:
            raise RuntimeError(f"Table 3.1 is missing {label}")
        points = []
        for index, year in year_columns:
            value = number_text(row[index] if index < len(row) else None)
            if value is None:
                continue
            points.append({"date": year, "value": value, "locator": {"row_label": label, "column": year}})
        series.append(
            _series(
                series_id,
                label,
                "millions of dollars",
                "nominal",
                "hist03z1",
                url,
                points,
                "Estimate columns in the source table are omitted.",
            )
        )
    gdp_names = {
        "national defense": "omb_fn_defense_pct_gdp",
        "human resources": "omb_fn_human_pct_gdp",
        "physical resources": "omb_fn_physical_pct_gdp",
        "net interest": "omb_fn_interest_pct_gdp",
        "other functions": "omb_fn_other_pct_gdp",
        "total, federal outlays": "omb_fn_total_pct_gdp",
    }
    for label, series_id in gdp_names.items():
        row = by_label.get(label + "|gdp")
        if row is None:
            raise RuntimeError(f"Table 3.1 percent of GDP is missing {label}")
        points = []
        published = _cell(row[0])
        for index, year in year_columns:
            value = number_text(row[index] if index < len(row) else None)
            if value is None:
                continue
            points.append({"date": year, "value": value, "locator": {"row_label": published, "column": year}})
        series.append(
            _series(
                series_id,
                published,
                "percent of GDP",
                "share_of_gdp",
                "hist03z1",
                url,
                points,
                "Estimate columns in the source table are omitted.",
            )
        )

    rows = _rows(paths["hist07z1"])
    title = _cell(rows[0][0])
    sources[4]["title"] = title
    url = BASE + "hist07z1_fy2027.xlsx"
    # Column positions are checked against the header text before use.
    if "Gross" not in _cell(rows[2][1]) or "Public" not in _cell(rows[2][3]):
        raise RuntimeError("Table 7.1 column labels did not match the expected debt layout")
    if "Percentages of GDP" not in _cell(rows[1][6]) or "Public" not in _cell(rows[2][8]):
        raise RuntimeError("Table 7.1 percent-of-GDP columns were not where expected")
    debt_cols = [
        ("omb_debt_gross", "Gross Federal Debt", 1, "millions of dollars", "nominal"),
        ("omb_debt_public", "Held by the Public, Total", 3, "millions of dollars", "nominal"),
        ("omb_debt_gross_pct_gdp", "Gross Federal Debt", 6, "percent of GDP", "share_of_gdp"),
        ("omb_debt_public_pct_gdp", "Held by the Public, Total", 8, "percent of GDP", "share_of_gdp"),
    ]
    for series_id, name, index, units, kind in debt_cols:
        series.append(
            _series(
                series_id,
                name,
                units,
                kind,
                "hist07z1",
                url,
                _year_points(rows, index, name),
                title,
            )
        )
    return sources, series
